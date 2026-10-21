# Helpers: python-pptx + raw OOXML (Morph, 3D models, gradients, text effects)
import copy, os
from lxml import etree
from PIL import Image, ImageFont
from pptx.util import Emu, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI
from pptx.dml.color import RGBColor

W, H = 12192000, 6858000
IN = 914400
NS = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
      'p': 'http://schemas.openxmlformats.org/presentationml/2006/main',
      'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
NSDECL = ' '.join(f'xmlns:{k}="{v}"' for k, v in NS.items())
MC = 'http://schemas.openxmlformats.org/markup-compatibility/2006'
RT_MODEL3D = 'http://schemas.microsoft.com/office/2017/06/relationships/model3d'

# --- charte Marin'Or Ivoire (mémoire p. 40) + two derived shades (dark/light backgrounds)
BORD, ROUGE, OR, JAUNE, CREME, ROSE = '7B1421', 'A40C21', 'F9A825', 'FBD021', 'FFECD1', 'F6E8EA'
VERT, VERTV, NOIR, BLANC = '2A5519', '376F23', '000000', 'FFFFFF'
DARK, LIGHT = '2A070D', 'FFF6EA'          # Bordeaux Profond darkened / Crème Douce lightened
F_TITLE, F_MED, F_SEMI, F_REG, F_XB = 'Cream Heavy', 'Montserrat Medium', 'Montserrat SemiBold', 'Montserrat', 'Montserrat ExtraBold'
FONTFILES = {F_TITLE: 'Cream-Heavy.ttf', F_MED: 'Montserrat-Medium.ttf', F_SEMI: 'Montserrat-SemiBold.ttf',
             F_REG: 'Montserrat-Regular.ttf', F_XB: 'Montserrat-ExtraBold.ttf', 'Montserrat Bold': 'Montserrat-Bold.ttf'}
TEXT_LOG = []   # (slide, name, paragraphs[(text, font, size, spc)], width_emu) for the 6-line check


def X(s):
    return etree.fromstring(f'<root {NSDECL}>{s}</root>')[0]


def spTree(slide):
    return slide.shapes._spTree


def next_id(slide):
    ids = [int(i) for i in spTree(slide).xpath('//@id') if str(i).isdigit()]
    return max(ids + [1]) + 1


def color_xml(c, alpha=None):
    a = f'<a:alpha val="{int(alpha * 1000)}"/>' if alpha is not None else ''
    return f'<a:srgbClr val="{c}">{a}</a:srgbClr>'


def set_fill(shape, xml_fill):
    spPr = shape._element.spPr
    for tag in ('a:noFill', 'a:solidFill', 'a:gradFill', 'a:blipFill', 'a:pattFill'):
        for e in spPr.findall(tag, NS):
            spPr.remove(e)
    geom = spPr.find('a:prstGeom', NS)
    el = X(xml_fill)
    if geom is not None:
        geom.addnext(el)
    else:
        spPr.append(el)


def set_line(shape, color=None, w_pt=0.75, alpha=None):
    spPr = shape._element.spPr
    for e in spPr.findall('a:ln', NS):
        spPr.remove(e)
    ln = f'<a:ln w="{int(w_pt * 12700)}"><a:solidFill>{color_xml(color, alpha)}</a:solidFill></a:ln>' if color else '<a:ln><a:noFill/></a:ln>'
    el = X(ln)
    fills = [e for e in spPr if etree.QName(e).localname in ('noFill', 'solidFill', 'gradFill', 'blipFill', 'prstGeom')]
    (fills[-1] if fills else spPr[-1]).addnext(el)


def add_effects(shape, shadow=None, soft=None, glow=None):
    """shadow=(blur_pt, dist_pt, alpha%) ; soft=edge radius pt ; glow=(rad_pt, color, alpha%)"""
    spPr = shape._element.spPr
    for e in spPr.findall('a:effectLst', NS):
        spPr.remove(e)
    parts = ''
    if glow:
        parts += f'<a:glow rad="{int(glow[0] * 12700)}">{color_xml(glow[1], glow[2])}</a:glow>'
    if shadow:
        parts += (f'<a:outerShdw blurRad="{int(shadow[0] * 12700)}" dist="{int(shadow[1] * 12700)}" dir="5400000" algn="t" rotWithShape="0">'
                  f'{color_xml("1A0507", shadow[2])}</a:outerShdw>')
    if soft:
        parts += f'<a:softEdge rad="{int(soft * 12700)}"/>'
    el = X(f'<a:effectLst>{parts}</a:effectLst>')
    ln = spPr.find('a:ln', NS)
    (ln if ln is not None else [e for e in spPr if etree.QName(e).localname in ('prstGeom', 'solidFill', 'gradFill', 'noFill', 'blipFill')][-1]).addnext(el)


def grad_xml(stops, angle=None, path=None):
    """stops: [(pos 0-100, color, alpha%|None)]"""
    gs = ''.join(f'<a:gs pos="{int(p * 1000)}">{color_xml(c, a)}</a:gs>' for p, c, a in stops)
    if path:
        l, t, r, b = path
        shade = f'<a:path path="circle"><a:fillToRect l="{l}" t="{t}" r="{r}" b="{b}"/></a:path>'
    else:
        shade = f'<a:lin ang="{int((angle or 0) * 60000)}" scaled="0"/>'
    return f'<a:gradFill rotWithShape="1"><a:gsLst>{gs}</a:gsLst>{shade}</a:gradFill>'


def rect(slide, name, x, y, w, h, fill=None, alpha=None, grad=None, line=None, line_w=0.75, line_alpha=None,
         geom=MSO_SHAPE.RECTANGLE, radius=None, shadow=None, soft=None):
    s = slide.shapes.add_shape(geom, Emu(int(x)), Emu(int(y)), Emu(int(w)), Emu(int(h)))
    s.name = name
    if grad:
        set_fill(s, grad)
    elif fill:
        set_fill(s, f'<a:solidFill>{color_xml(fill, alpha)}</a:solidFill>')
    else:
        set_fill(s, '<a:noFill/>')
    set_line(s, line, line_w, line_alpha)
    if radius is not None and geom == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    if shadow or soft:
        add_effects(s, shadow=shadow, soft=soft)
    # no text style inheritance surprises
    s.text_frame.text = ''
    st = s._element.find('p:style', NS)
    if st is not None:
        s._element.remove(st)
    return s


def picture(slide, name, path, x, y, w, h, fit='fill', radius=None, shadow=None, alpha=None, soft=None, descr=''):
    """fit='fill' crops the image to the box; fit='contain' keeps the full image inside (box shrinks)."""
    iw, ih = Image.open(path).size
    if fit == 'contain':
        r = min(w / iw, h / ih); nw, nh = iw * r, ih * r
        x, y, w, h = x + (w - nw) / 2, y + (h - nh) / 2, nw, nh
    p = slide.shapes.add_picture(path, Emu(int(x)), Emu(int(y)), Emu(int(w)), Emu(int(h)))
    p.name = name
    if descr:
        p._element.nvPicPr.cNvPr.set('descr', descr)
    if fit == 'fill':
        br, ir = w / h, iw / ih
        if ir > br:
            c = (1 - br / ir) / 2; p.crop_left = c; p.crop_right = c
        else:
            c = (1 - ir / br) / 2; p.crop_top = c; p.crop_bottom = c
    if radius:
        p.auto_shape_type = MSO_SHAPE.ROUNDED_RECTANGLE
        geom = p._element.spPr.find('a:prstGeom', NS)
        av = geom.find('a:avLst', NS)
        av.append(X(f'<a:gd name="adj" fmla="val {int(radius * 100000)}"/>'))
    if alpha is not None:
        blip = p._element.blipFill.find('a:blip', NS)
        blip.append(X(f'<a:alphaModFix amt="{int(alpha * 1000)}"/>'))
    if shadow or soft:
        add_effects(p, shadow=shadow, soft=soft)
    return p


def text(slide, sn, name, x, y, w, h, paras, align='l', anchor='t', line_sp=None, wrap=True, autofit=False, row=None):
    """paras: list of paragraphs; each paragraph = list of runs (txt, font, size, color, dict(opts))
       opts: spc (1/100 pt), alpha, grad (list of stops), outline (color, w_pt), bold, caps, space_after"""
    tb = slide.shapes.add_textbox(Emu(int(x)), Emu(int(y)), Emu(int(w)), Emu(int(h)))
    tb.name = name
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {'t': MSO_ANCHOR.TOP, 'm': MSO_ANCHOR.MIDDLE, 'b': MSO_ANCHOR.BOTTOM}[anchor]
    bodyPr = tf._txBody.find('a:bodyPr', NS)
    for e in list(bodyPr):
        bodyPr.remove(e)
    bodyPr.append(X('<a:normAutofit/>' if autofit else '<a:noAutofit/>'))
    log = []
    for i, runs in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {'l': PP_ALIGN.LEFT, 'c': PP_ALIGN.CENTER, 'r': PP_ALIGN.RIGHT}[align]
        if line_sp:
            p.line_spacing = line_sp
        plog = []
        for (t, font, size, col, *o) in runs:
            o = o[0] if o else {}
            r = p.add_run(); r.text = t
            f = r.font; f.name = font; f.size = Pt(size)
            if o.get('bold'):
                f.bold = True
            rPr = r._r.get_or_add_rPr()
            if o.get('spc'):
                rPr.set('spc', str(o['spc']))
            if o.get('caps'):
                rPr.set('cap', 'all')
            for e in rPr.findall('a:solidFill', NS):
                rPr.remove(e)
            fill = grad_xml(o['grad'], angle=90) if o.get('grad') else f'<a:solidFill>{color_xml(col, o.get("alpha"))}</a:solidFill>'
            ln = ''
            if o.get('outline'):
                oc, ow = o['outline']
                ln = f'<a:ln w="{int(ow * 12700)}"><a:solidFill>{color_xml(oc, o.get("outline_alpha"))}</a:solidFill></a:ln>'
                if o.get('hollow'):
                    fill = '<a:noFill/>'
            # order inside rPr: ln, fill, effectLst, latin...
            latin = rPr.find('a:latin', NS)
            anchor_el = latin
            if ln:
                el = X(ln); (anchor_el.addprevious(el) if anchor_el is not None else rPr.append(el))
            el = X(fill); (anchor_el.addprevious(el) if anchor_el is not None else rPr.append(el))
            if o.get('glow'):
                g = o['glow']
                el = X(f'<a:effectLst><a:glow rad="{int(g[0] * 12700)}">{color_xml(g[1], g[2])}</a:glow></a:effectLst>')
                anchor_el.addprevious(el)
            if latin is not None:
                latin.set('typeface', font)
                for tag in ('a:ea', 'a:cs'):
                    pass
            plog.append((t, font, size, o.get('spc', 0), o.get('caps', False)))
        if runs and runs[0][4:] and runs[0][4].get('space_after'):
            p.space_after = Pt(runs[0][4]['space_after'])
        log.append(plog)
    TEXT_LOG.append((sn, name, log, w, row))
    return tb


# ---------- Morph transition
def morph(slide, dur=1400, option='byObject'):
    sld = slide._element
    for e in sld.findall('p:transition', NS):
        sld.remove(e)
    for e in sld.findall(f'{{{MC}}}AlternateContent'):
        if e.find(f'.//{{{NS["p"]}}}transition') is not None:
            sld.remove(e)
    xml = (f'<mc:AlternateContent xmlns:mc="{MC}"><mc:Choice xmlns:p159="http://schemas.microsoft.com/office/powerpoint/2015/09/main" Requires="p159">'
           f'<p:transition xmlns:p14="http://schemas.microsoft.com/office/powerpoint/2010/main" spd="slow" p14:dur="{dur}"><p159:morph option="{option}"/></p:transition>'
           f'</mc:Choice><mc:Fallback><p:transition spd="slow"><p:fade/></p:transition></mc:Fallback></mc:AlternateContent>')
    el = X(xml)
    clr = sld.find('p:clrMapOvr', NS)
    clr.addnext(el)


# ---------- 3D model (Office 2019 am3d), shared GLB part per product
_GLB = {}


def glb_part(prs, key, path):
    if key not in _GLB:
        n = len(_GLB) + 1
        _GLB[key] = Part(PackURI(f'/ppt/media/model3d{n}.glb'), 'model/gltf-binary', prs.part.package, open(path, 'rb').read())
    return _GLB[key]


LIGHTS = ('<am3d:ambientLight><am3d:clr><a:scrgbClr r="50000" g="50000" b="50000"/></am3d:clr><am3d:illuminance n="500000" d="1000000"/></am3d:ambientLight>'
          '<am3d:ptLight rad="0"><am3d:clr><a:scrgbClr r="100000" g="75000" b="50000"/></am3d:clr><am3d:intensity n="9765625" d="1000000"/><am3d:pos x="21959998" y="70920001" z="16344003"/></am3d:ptLight>'
          '<am3d:ptLight rad="0"><am3d:clr><a:scrgbClr r="40000" g="60000" b="95000"/></am3d:clr><am3d:intensity n="12250000" d="1000000"/><am3d:pos x="-37964106" y="51130435" z="57631972"/></am3d:ptLight>'
          '<am3d:ptLight rad="0"><am3d:clr><a:scrgbClr r="86837" g="72700" b="100000"/></am3d:clr><am3d:intensity n="3125000" d="1000000"/><am3d:pos x="-37739122" y="58056624" z="-34769649"/></am3d:ptLight>')


def model3d(prs, slide, inst, glb_path, name, descr):
    rid_glb = slide.part.relate_to(glb_part(prs, inst['key'], glb_path), RT_MODEL3D)
    _, rid_png = slide.part.get_or_add_image_part(inst['png'])
    sid = next_id(slide)
    (x, y), (cx, cy) = inst['off'], inst['ext']
    L = inst['L']
    mpmu = 1.0 / L                                   # model normalised to 1 m on its largest side
    pre = [int(-c * mpmu * 36000000) for c in inst['center']]
    ang = lambda d: int(round(d * 60000)) % 21600000
    descr = descr.replace('"', '&quot;').replace("'", '&apos;')
    xml = f'''<mc:AlternateContent xmlns:mc="{MC}"><mc:Choice xmlns:am3d="http://schemas.microsoft.com/office/drawing/2017/model3d" Requires="am3d">
<p:graphicFrame><p:nvGraphicFramePr><p:cNvPr id="{sid}" name="{name}" descr="{descr}"/><p:cNvGraphicFramePr><a:graphicFrameLocks noChangeAspect="1"/></p:cNvGraphicFramePr><p:nvPr/></p:nvGraphicFramePr>
<p:xfrm><a:off x="{x}" y="{y}"/><a:ext cx="{cx}" cy="{cy}"/></p:xfrm>
<a:graphic><a:graphicData uri="http://schemas.microsoft.com/office/drawing/2017/model3d">
<am3d:model3d r:embed="{rid_glb}">
<am3d:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="{cx}" cy="{cy}"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></am3d:spPr>
<am3d:camera><am3d:pos x="0" y="0" z="67740115"/><am3d:up dx="0" dy="36000000" dz="0"/><am3d:lookAt x="0" y="0" z="0"/><am3d:perspective fov="2700000"/></am3d:camera>
<am3d:trans><am3d:meterPerModelUnit n="{int(round(mpmu * 1000000))}" d="1000000"/><am3d:preTrans dx="{pre[0]}" dy="{pre[1]}" dz="{pre[2]}"/>
<am3d:scale><am3d:sx n="1000000" d="1000000"/><am3d:sy n="1000000" d="1000000"/><am3d:sz n="1000000" d="1000000"/></am3d:scale>
<am3d:rot ax="{ang(inst['ax'])}" ay="{ang(inst['ay'])}" az="{ang(inst['az'])}"/><am3d:postTrans dx="0" dy="0" dz="0"/></am3d:trans>
<am3d:raster rName="Office3DRenderer" rVer="16.0.8326"><am3d:blip r:embed="{rid_png}"/></am3d:raster>
<am3d:objViewport viewportSz="{inst['V']}"/>{LIGHTS}
</am3d:model3d></a:graphicData></a:graphic></p:graphicFrame></mc:Choice>
<mc:Fallback><p:pic><p:nvPicPr><p:cNvPr id="{sid}" name="{name}" descr="{descr}"/><p:cNvPicPr><a:picLocks noGrp="1" noRot="1" noChangeAspect="1" noMove="1" noResize="1" noEditPoints="1" noAdjustHandles="1" noChangeArrowheads="1" noChangeShapeType="1" noCrop="1"/></p:cNvPicPr><p:nvPr/></p:nvPicPr>
<p:blipFill><a:blip r:embed="{rid_png}"/><a:stretch><a:fillRect/></a:stretch></p:blipFill>
<p:spPr><a:xfrm><a:off x="{x}" y="{y}"/><a:ext cx="{cx}" cy="{cy}"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr></p:pic></mc:Fallback></mc:AlternateContent>'''
    spTree(slide).append(X(xml))


def notes(slide, txt):
    slide.notes_slide.notes_text_frame.text = txt


# ---------- 6-line control: measure wrapped lines with the real fonts
_FC = {}


def _font(name, size):
    k = (name, size)
    if k not in _FC:
        _FC[k] = ImageFont.truetype(os.path.expanduser('~/.fonts/' + FONTFILES.get(name, 'Montserrat-Medium.ttf')), int(size * 10))
    return _FC[k]


def count_lines(paras, width_emu):
    width_pt = width_emu / 12700
    total = 0
    for plog in paras:
        words = []
        for (t, font, size, spc, caps) in plog:
            for w in (t.upper() if caps else t).replace('\n', ' \n ').split(' '):
                words.append((w, font, size, spc))
        lines, cur = 1, 0.0
        for (w, font, size, spc) in words:
            if w == '\n':
                lines += 1; cur = 0; continue
            f = _font(font, size)
            ww = f.getlength(w + ' ') / 10 + len(w) * spc / 100
            if cur + ww > width_pt + 2 and cur > 0:
                lines += 1; cur = ww
            else:
                cur += ww
        total += lines if any(t.strip() for (t, *_) in plog) else 0
    return total


# ---------- entrance animations (PowerPoint 2019+ timing XML, validated by the Open XML SDK)
RISE = 20 / 720          # 20 px on a 1280 x 720 px slide, as a fraction of the slide height (ppt_y units)


def _eff(ids, kind, spid, node, dur):
    """One entrance effect. kind: float (fade + 20 px rise), wipe (from left), wheel, fade. Ease-out = decel 100 %."""
    a, b, c, d = next(ids), next(ids), next(ids), next(ids)
    tgt = f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>'
    vis = (f'<p:set><p:cBhvr><p:cTn id="{b}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>{tgt}'
           f'<p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal val="visible"/></p:to></p:set>')
    preset = {'float': (42, 0), 'wipe': (22, 8), 'wheel': (21, 1), 'fade': (10, 0)}[kind]
    filt = {'float': 'fade', 'wipe': 'wipe(left)', 'wheel': 'wheel(1)', 'fade': 'fade'}[kind]
    body = vis + (f'<p:animEffect transition="in" filter="{filt}"><p:cBhvr><p:cTn id="{c}" dur="{dur}" decel="100000"/>{tgt}</p:cBhvr></p:animEffect>')
    if kind == 'float':
        e = next(ids)
        body += (f'<p:anim calcmode="lin" valueType="num"><p:cBhvr><p:cTn id="{d}" dur="{dur}" decel="100000" fill="hold"/>{tgt}'
                 f'<p:attrNameLst><p:attrName>ppt_x</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst>'
                 f'<p:tav tm="0"><p:val><p:strVal val="#ppt_x"/></p:val></p:tav><p:tav tm="100000"><p:val><p:strVal val="#ppt_x"/></p:val></p:tav></p:tavLst></p:anim>'
                 f'<p:anim calcmode="lin" valueType="num"><p:cBhvr><p:cTn id="{e}" dur="{dur}" decel="100000" fill="hold"/>{tgt}'
                 f'<p:attrNameLst><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst>'
                 f'<p:tav tm="0"><p:val><p:strVal val="#ppt_y+{RISE:.4f}"/></p:val></p:tav><p:tav tm="100000"><p:val><p:strVal val="#ppt_y"/></p:val></p:tav></p:tavLst></p:anim>')
    return (f'<p:par><p:cTn id="{a}" presetID="{preset[0]}" presetClass="entr" presetSubtype="{preset[1]}" fill="hold" grpId="0" nodeType="{node}">'
            f'<p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>{body}</p:childTnLst></p:cTn></p:par>')


def build_timing(slide, auto_steps, click_groups, bld):
    """auto_steps: steps run at slide start; click_groups: list of groups, each a list of steps.
       step = list of (kind, spid, dur_ms). Steps of a group run one after the other ("après la précédente")."""
    sld = slide._element
    old = sld.find('p:timing', NS)
    media = []
    if old is not None:                      # keep python-pptx video media nodes
        media = old.findall('.//p:video', NS) + old.findall('.//p:audio', NS)
        sld.remove(old)
    if not auto_steps and not click_groups and not media:
        return
    cnt = iter(range(3, 100000))
    ids = cnt
    out = ''

    def group(steps, auto):
        nonlocal out
        gid = next(ids)
        cond = ('<p:cond delay="indefinite"/><p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond>' if auto else '<p:cond delay="indefinite"/>')
        inner, t = '', 0
        for k, step in enumerate(steps):
            sid = next(ids)
            effs = ''
            for j, (kind, spid, dur) in enumerate(step):
                if k == 0 and j == 0:
                    node = 'afterEffect' if auto else 'clickEffect'
                elif j == 0:
                    node = 'afterEffect'
                else:
                    node = 'withEffect'
                effs += _eff(ids, kind, spid, node, dur)
            inner += f'<p:par><p:cTn id="{sid}" fill="hold"><p:stCondLst><p:cond delay="{t}"/></p:stCondLst><p:childTnLst>{effs}</p:childTnLst></p:cTn></p:par>'
            t += max(d for (_, _, d) in step)
        out += f'<p:par><p:cTn id="{gid}" fill="hold"><p:stCondLst>{cond}</p:stCondLst><p:childTnLst>{inner}</p:childTnLst></p:cTn></p:par>'

    if auto_steps:
        group(auto_steps, True)
    for g in click_groups:
        group(g, False)
    seq = ''
    if out:
        seq = (f'<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>{out}</p:childTnLst></p:cTn>'
               '<p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
               '<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>')
    bl = ''.join(f'<p:bldP spid="{spid}" grpId="0"{" animBg=\"1\"" if bg else ""}/>' for spid, bg in bld)
    xml = (f'<p:timing><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>{seq}</p:childTnLst></p:cTn></p:par></p:tnLst>'
           + (f'<p:bldLst>{bl}</p:bldLst>' if bl else '') + '</p:timing>')
    tm = X(xml)
    root_list = tm.find('.//p:childTnLst', NS)
    for mnode in media:                      # renumber media cTn ids after ours
        for c in mnode.iter(f'{{{NS["p"]}}}cTn'):
            c.set('id', str(next(ids)))
        root_list.append(mnode)
    # p:timing goes after the transition (or clrMapOvr) and before extLst
    ext = sld.find('p:extLst', NS)
    if ext is not None:
        ext.addprevious(tm)
    else:
        sld.append(tm)
