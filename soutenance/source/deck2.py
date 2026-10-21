# Soutenance Marin'Or Ivoire, version 2 : 46 slides, fonds unis, Morph réservé aux produits 3D,
# une seule règle d'apparition (fondu + 20 px vers le haut, 0,5 s, ease-out), graphiques sur mesure.
# Internal numbering = v1 numbering (2..47); displayed number = i - 1.
import json, math, os, sys, re
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.enum.shapes import MSO_SHAPE
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from lib import *
import lib
from layout3d import NAMES, WID

B = os.path.dirname(os.path.abspath(__file__))
A = B + '/assets/'
VID = B + '/video/'
INST = {int(k): v for k, v in json.load(open(B + '/instances.json')).items()}
OUT = sys.argv[1] if len(sys.argv) > 1 else B + '/out/Soutenance_MarinOr_Ivoire.pptx'
NO3D = bool(os.environ.get('NO3D'))
I = lambda v: int(v * IN)

prs = Presentation()
prs.slide_width, prs.slide_height = Emu(W), Emu(H)
BLANK = prs.slide_layouts[6]
SL = {}
PRODUCT_SLIDES = range(13, 19)          # Morph 3D sequence (displayed 12 to 17)

# Palette logic inspired by the reference deck: one deep brand colour as solid background for key slides,
# white for content, a very light brand tint for cards, a tinted dark text + a muted secondary, gold sparingly.
TXT, MUTED, TINT = '2A070D', '6E5A5D', 'D9AEB5'      # derived from Bordeaux Profond (signalled in the notes)
BORD_HI, BORD_LO = '8A1828', '68101C'                # subtle "smart" gradient on dark slides


def C(mode):
    if mode == 'D':
        return dict(title=BLANC, body=CREME, kick=OR, muted='E8C9CC')
    return dict(title=BORD, body=TXT, kick=MUTED, muted=MUTED)


# ---------------------------------------------------------------- slide scaffolding
def new(i, mode, solid=False, dur=None):
    s = prs.slides.add_slide(BLANK)
    SL[i] = dict(slide=s, mode=mode, morph=dur, groups={}, chart=None)
    if mode == 'D':
        if solid:
            rect(s, '!!BG', 0, 0, W, H, fill=BORD)
        else:
            rect(s, '!!BG', 0, 0, W, H, grad=grad_xml([(0, BORD_HI, None), (100, BORD_LO, None)], angle=35))
    else:
        rect(s, '!!BG', 0, 0, W, H, fill=BLANC)
    return s


def ag(i, g, *shapes):
    """assign shapes to an animation group (0 = automatic with the slide, 1.. = click groups)"""
    for sh in shapes:
        SL[i]['groups'][sh.shape_id] = (g, sh)


def finish(s, i, mode, kicker=None, logo=True, note='', secs=20):
    c = C(mode)
    if kicker:
        text(s, i, '!!Kicker', I(0.75), I(0.52), I(8.6), I(0.3), [[(kicker, F_SEMI, 10.5, c['kick'], dict(spc=220, caps=True))]], row='kicker')
    text(s, i, '!!PageNo', W - I(1.35), H - I(0.6), I(0.6), I(0.25), [[(f'{i - 1:02d}', F_SEMI, 10, c['kick'], dict(spc=150))]], align='r', row='chrome')
    if logo:
        lw = I(1.05); lh = lw * 669 / 1274
        picture(s, '!!Logo', A + ('logo_blanc.png' if mode == 'D' else 'logo_couleur.png'), W - I(0.75) - lw, I(0.34), lw, lh,
                fit='contain', descr="Logo Marin'Or Ivoire")
    SL[i]['secs'] = secs
    notes(s, f'[{secs} s] ' + note)


def title(s, i, mode, txt, x, y, w, h, size=40, align='l', anchor='t', auto=True):
    t = text(s, i, '!!Title', x, y, w, h, [[(txt, F_TITLE, size, C(mode)['title'], {})]], align=align, anchor=anchor, line_sp=0.95)
    if auto and i not in PRODUCT_SLIDES:
        ag(i, 0, t)
    return t


def body(s, i, mode, name, paras, x, y, w, h, size=18, align='l', line_sp=1.12, color=None, row=None, font=F_MED, space=8):
    c = C(mode)
    P = []
    for p in paras:
        if isinstance(p, str):
            P.append([(p, font, size, color or c['body'], dict(space_after=space))])
        else:
            P.append([(t, f or font, sz or size, col or color or c['body'], dict(o, space_after=space)) for (t, f, sz, col, o) in p])
    return text(s, i, name, x, y, w, h, P, align=align, line_sp=line_sp, row=row)


def card(s, name, x, y, w, h, fill=ROSE):
    return rect(s, name, x, y, w, h, fill=fill, geom=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.07)


def pic(s, name, path, x, y, w, h, **k):
    k.setdefault('shadow', (14, 3, 14))
    return picture(s, name, path, x, y, w, h, **k)


def badge(s, i, name, n, x, y, d=I(0.46), mode='L'):
    b = rect(s, name, x, y, d, d, fill=BORD if mode == 'L' else OR, geom=MSO_SHAPE.OVAL)
    t = text(s, None, name + 'T', x, y, d, d, [[(n, F_SEMI, 12.5, BLANC if mode == 'L' else BORD, {})]], align='c', anchor='m')
    return b, t


def products(s, i, only=None):
    for inst in INST[i]:
        if only and inst['key'] not in only:
            continue
        if inst['shadow']:
            sw = WID[inst['key']] * inst['k'] * 1.25; sh = sw * 0.2
            rect(s, f"!!S_{inst['key']}", inst['foot'][0] - sw / 2, inst['foot'][1] - sh * 0.55, sw, sh, geom=MSO_SHAPE.OVAL,
                 grad=grad_xml([(0, '1A0507', 38), (100, '1A0507', 0)], path=(50000, 50000, 50000, 50000)))
        if NO3D:
            (x, y), (cx, cy) = inst['off'], inst['ext']
            picture(s, f"!!P_{inst['key']}", inst['png'], x, y, cx, cy, fit='contain', descr=NAMES[inst['key']])
        else:
            model3d(prs, s, inst, f"{B}/p3d/{inst['key']}.glb", f"!!P_{inst['key']}", f"Modèle 3D : {NAMES[inst['key']]}")


L0, R0 = I(0.75), W - I(0.75)
TW = R0 - L0

# ================================================================ SLIDES (v1 slide 1 « Bienvenue » removed)
# 2 — Page de garde
i = 2; m = 'L'; s = new(i, m)
h = pic(s, '!!Hero', A + 'packshot.jpg', I(9.2), I(1.25), I(3.38), I(4.25), radius=0.05, descr="Packshot : Graines d'akpi, Huile rouge, Poudre de crevettes")
title(s, i, m, "Conception et réalisation de supports de communication pour la promotion de Marin'Or Ivoire", L0, I(1.75), I(8.2), I(3.0), size=30)
ag(i, 1, h)
for k, (lab, nm) in enumerate([('Présenté par', 'Ekien Kouadio Constant Ephrem'), ('Sous la direction de', 'Dr Tayoro Gnewa Gérard'),
                               ('Sous la supervision de', 'TOA Agnini Jules Evariste, Professeur Titulaire')]):
    x = L0 + [0, I(3.35), I(6.7)][k]; cw = [I(3.2), I(3.2), I(5.1)][k]
    r = rect(s, f'!!Rule{k}', x, I(5.72), I(0.5), I(0.03), fill=OR)
    a = text(s, i, f'!!Lab{k}', x, I(5.86), cw, I(0.25), [[(lab, F_SEMI, 9.5, BORD, dict(spc=150, caps=True))]], row='lab')
    b = text(s, i, f'!!Nm{k}', x, I(6.13), cw, I(0.6), [[(nm, F_MED, 12.5, TXT)]], row='nm')
    ag(i, 1, r, a, b)
finish(s, i, m, 'ISTC Polytechnique · Année académique 2025-2026', secs=25,
       note="Bonjour. Mesdames et Messieurs les membres du jury, merci de votre présence. Mon mémoire s'intitule : « Conception et réalisation de supports de communication pour la promotion de Marin'Or Ivoire ». Il a été réalisé sous la direction du Dr Tayoro Gnewa Gérard et sous la supervision de M. TOA Agnini Jules Evariste, Professeur Titulaire.")

# 3 — Plan
PLAN = [('01', 'Introduction et problématique'), ('02', 'Les produits'), ('03', 'Identité visuelle'),
        ('04', 'Supports print et audiovisuels'), ('05', 'Le site web'), ('06', 'Bilan et conclusion')]
i = 3; m = 'L'; s = new(i, m)
title(s, i, m, 'Plan de la présentation', L0, I(1.0), I(8), I(0.8), size=40)
for k, (n, t) in enumerate(PLAN):
    col, r = k // 3, k % 3
    x = L0 + col * I(5.7); y = I(2.55) + r * I(1.25)
    b, bt = badge(s, i, f'!!PlanB{k}', n, x, y + I(0.22))
    ln = rect(s, f'!!PlanLine{k}', x, y + I(0.98), I(5.0), I(0.012), fill=BORD, alpha=18)
    tt = text(s, i, f'!!PlanItem{k}', x + I(0.68), y, I(4.6), I(0.9), [[(t, F_SEMI, 18, TXT, {})]], anchor='m', row=f'plan{r}')
    ag(i, 1, b, bt, ln, tt)
finish(s, i, m, 'Sommaire', secs=15, note="Mon exposé suit six temps : l'introduction et la problématique, les produits, l'identité visuelle, les supports print et audiovisuels, le site web, puis le bilan.")

# 4 — Contexte
i = 4; m = 'L'; s = new(i, m)
h = pic(s, '!!Hero', A + 'food.jpg', I(7.35), I(1.0), I(5.25), I(5.55), radius=0.05, descr='Plat traditionnel partagé')
title(s, i, m, 'Contexte', L0, I(1.0), I(6.2), I(0.8), size=40)
b = body(s, i, m, '!!Body', ["Les consommateurs recherchent des marques qui valorisent leur patrimoine culinaire.",
                             "De jeunes marques locales peinent pourtant à s'imposer, faute de moyens ou de structuration."],
         L0, I(2.3), I(6.3), I(3.2), size=18, space=14)
ag(i, 1, b, h)
finish(s, i, m, '01 · Introduction', secs=20,
       note="Dans nos sociétés, la recherche de saveurs authentiques grandit. Les consommateurs se tournent vers des marques capables de valoriser leur patrimoine culinaire. Pourtant, nombre de jeunes marques locales peinent à se faire une place, faute de moyens ou de structuration (Introduction, p. 16).")

# 5 — Marin'Or Ivoire
i = 5; m = 'L'; s = new(i, m)
p = card(s, '!!Panel', I(7.3), I(1.25), I(5.3), I(4.9))
lw = I(4.4); lh = lw * 669 / 1274
lg = picture(s, '!!LogoHero', A + 'logo_couleur.png', I(7.3) + I(5.3) / 2 - lw / 2, I(1.25) + I(4.9) / 2 - lh / 2, lw, lh, fit='contain', descr="Logo Marin'Or Ivoire")
title(s, i, m, "Marin'Or Ivoire", L0, I(1.0), I(6.2), I(0.8), size=40)
b = body(s, i, m, '!!Body', ["Des marinades et épices authentiques, portées par M. Moustapha.",
                             "Point de départ : un stage au lancement de la marque."], L0, I(2.3), I(6.3), I(3.8), size=18, space=14)
ag(i, 1, p, lg, b)
finish(s, i, m, '01 · Introduction', logo=False, secs=20,
       note="Marin'Or Ivoire est un projet porté par M. Moustapha : des marinades et épices préparées, authentiques et accessibles, qui simplifient la cuisine du quotidien (p. 20). J'ai conçu ses premiers supports lors d'un stage ; ce travail restait ponctuel, sans stratégie globale.")

# 6 — Constats
CON = ['Une identité visuelle incomplète', 'Une visibilité restreinte', 'Aucune plateforme numérique', 'Un potentiel sous-exploité']
i = 6; m = 'L'; s = new(i, m)
title(s, i, m, "Constats de l'existant", L0, I(1.0), I(9), I(0.8), size=40)
for k, t in enumerate(CON):
    x = L0 + (k % 2) * I(5.95); y = I(2.45) + (k // 2) * I(1.75)
    cd = card(s, f'!!Card{k}', x, y, I(5.65), I(1.45))
    b, bt = badge(s, i, f'!!CardB{k}', f'0{k + 1}', x + I(0.35), y + I(0.495))
    tt = text(s, i, f'!!CardT{k}', x + I(1.0), y, I(4.4), I(1.45), [[(t, F_SEMI, 18, TXT, {})]], anchor='m', row=f'card{k // 2}')
    ag(i, k + 1, cd, b, bt, tt)
finish(s, i, m, '01 · Analyse de l’existant', secs=25,
       note="L'analyse fait ressortir quatre constats (p. 25) : une identité visuelle incomplète, sans charte formalisée ; une visibilité restreinte ; aucun site web ; un potentiel sous-exploité. La communication est assurée par le porteur du projet lui-même, avec des moyens limités.")

# 7 — Question centrale (fond bordeaux, dégradé subtil)
i = 7; m = 'D'; s = new(i, m)
q = [[("Comment concevoir et réaliser des supports de communication efficaces, permettant à Marin'Or Ivoire de construire une ", F_TITLE, 29, BLANC, {}),
      ("identité de marque forte et durable", F_TITLE, 29, OR, {}),
      (" dans le contexte concurrentiel ivoirien ?", F_TITLE, 29, BLANC, {})]]
t = text(s, i, '!!Title', I(1.0), I(1.55), W - I(2.0), I(4.2), q, align='c', anchor='m', line_sp=1.05)
ag(i, 0, t)
finish(s, i, m, 'Problématique · question centrale', secs=25,
       note="D'où la question centrale (p. 17) : comment concevoir et réaliser des supports de communication efficaces, permettant à Marin'Or Ivoire de construire une identité de marque forte et durable dans le contexte concurrentiel ivoirien ?")

# 8 — Hypothèse
i = 8; m = 'D'; s = new(i, m)
title(s, i, m, 'Hypothèse', L0, I(1.0), I(8), I(0.8), size=40)
b = body(s, i, m, '!!Body', ["Une stratégie combinant supports imprimés, contenus audiovisuels et présence numérique permettrait à la marque de gagner en visibilité et en crédibilité."],
         L0, I(2.25), I(9.6), I(1.9), size=21)
ag(i, 1, b)
for k, t in enumerate(['Imprimé', 'Audiovisuel', 'Numérique']):
    x = L0 + k * I(2.75)
    pl = rect(s, f'!!Pill{k}', x, I(4.75), I(2.45), I(0.62), fill=BLANC, alpha=12, geom=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5, line=OR, line_w=1.0)
    pt = text(s, i, f'!!PillT{k}', x, I(4.75), I(2.45), I(0.62), [[(t, F_SEMI, 15, BLANC, dict(spc=100))]], align='c', anchor='m', row='pills')
    ag(i, 2, pl, pt)
finish(s, i, m, 'Problématique · hypothèse', secs=20,
       note="Nous formulons l'hypothèse qu'une stratégie combinant supports imprimés, contenus audiovisuels et présence numérique permettrait à la marque de gagner significativement en visibilité et en crédibilité auprès de son public cible (p. 17).")

# 9 — Méthodologie
METH = [('Recherche documentaire', 'Ouvrages, mémoires antérieurs'), ('Observation directe', 'Pendant le stage'),
        ('Entretiens semi-directifs', 'Porteur du projet, consommateurs')]
i = 9; m = 'L'; s = new(i, m)
title(s, i, m, 'Méthodologie qualitative', L0, I(1.0), I(9), I(0.8), size=40)
for k, (t, d) in enumerate(METH):
    x = L0 + k * I(3.95)
    cd = card(s, f'!!Card{k}', x, I(2.35), I(3.7), I(2.35))
    b, bt = badge(s, i, f'!!CardB{k}', f'0{k + 1}', x + I(0.3), I(2.65))
    tt = text(s, i, f'!!CardT{k}', x + I(0.3), I(3.3), I(3.2), I(0.45), [[(t, F_SEMI, 15.5, TXT, {})]], row='cardt')
    dd = text(s, i, f'!!CardD{k}', x + I(0.3), I(3.77), I(3.2), I(0.8), [[(d, F_MED, 12.5, MUTED, {})]], row='cardd')
    ag(i, k + 1, cd, b, bt, tt, dd)
f = body(s, i, m, '!!Foot', [[('Terrain : Abidjan', F_SEMI, 14, BORD, {}), (' · diagnostic empirique et analyse comparative des marques concurrentes', F_MED, 14, TXT, {})]],
         L0, I(5.2), TW, I(0.4), row='foot')
ag(i, 4, f)
finish(s, i, m, '01 · Méthodologie', secs=25,
       note="Démarche qualitative à Abidjan (p. 23) : recherche documentaire, observation directe pendant le stage, entretiens semi-directifs avec le porteur du projet et des consommateurs. Conformément aux objectifs d'une licence professionnelle, l'étude repose sur un diagnostic de terrain empirique et une analyse comparative des marques concurrentes, plutôt que sur une étude de marché quantitative. (Cette dernière phrase n'est pas dans le mémoire : à assumer à l'oral.)")

# 10 — Objectifs
OBJ = ['Accroître la visibilité', 'Construire une identité cohérente', 'Produire imprimé, audiovisuel et web', 'Poser les bases d’une présence numérique']
i = 10; m = 'L'; s = new(i, m)
title(s, i, m, 'Objectifs', L0, I(1.0), I(8), I(0.8), size=40)
b = body(s, i, m, '!!Body', ["Doter Marin'Or Ivoire d'une identité visuelle forte et de supports adaptés, pour renforcer sa notoriété."],
         L0, I(2.15), I(10.4), I(1.2), size=20)
ag(i, 1, b)
for k, t in enumerate(OBJ):
    x = L0 + (k % 2) * I(5.95); y = I(3.65) + (k // 2) * I(1.15)
    cd = card(s, f'!!Card{k}', x, y, I(5.65), I(0.92))
    bb, bt = badge(s, i, f'!!CardB{k}', f'0{k + 1}', x + I(0.3), y + I(0.23))
    tt = text(s, i, f'!!CardT{k}', x + I(0.95), y, I(4.55), I(0.92), [[(t, F_SEMI, 14, TXT, {})]], anchor='m', row=f'obj{k // 2}')
    ag(i, 2, cd, bb, bt, tt)
finish(s, i, m, '01 · Objectifs', secs=20,
       note="Objectif général (p. 21) : doter Marin'Or Ivoire d'une identité visuelle forte et de supports adaptés. Quatre objectifs spécifiques : accroître la visibilité, construire une identité cohérente, produire des supports imprimés, audiovisuels et numériques, poser les bases d'une présence numérique.")

# 11 — Public cible
i = 11; m = 'L'; s = new(i, m)
h = pic(s, '!!Hero', A + 'affiche2_situ.jpg', I(8.3), I(1.0), I(4.3), I(5.6), radius=0.05, descr='Affiche 2 en situation')
title(s, i, m, 'Public cible', L0, I(1.0), I(7), I(0.8), size=40)
bg_ = text(s, i, '!!Big', L0, I(2.05), I(7), I(1.25), [[('25 – 50 ans', F_TITLE, 66, ROUGE, {})]])
b = body(s, i, m, '!!Body', ["Les ménages d'Abidjan qui cuisinent régulièrement pour leur foyer."], L0, I(3.45), I(7.0), I(1.0), size=19)
b2 = body(s, i, m, '!!Body2', ['Aussi : distributeurs, restaurateurs, passionnés de cuisine.'], L0, I(4.55), I(7.0), I(0.6), size=15, color=MUTED)
ag(i, 1, bg_, b, h); ag(i, 2, b2)
finish(s, i, m, '01 · Cahier des charges', secs=20,
       note="La cible prioritaire (p. 27) : les ménages d'Abidjan, femmes et hommes de 25 à 50 ans qui cuisinent régulièrement. Plus largement : consommateurs sensibles aux produits locaux, distributeurs, restaurateurs et passionnés de cuisine.")

# 12 — Axes
AX = ["L'authenticité des saveurs ivoiriennes", 'La convivialité et le partage', 'La simplicité accessible au quotidien']
i = 12; m = 'L'; s = new(i, m)
title(s, i, m, 'Trois axes de communication', L0, I(1.0), I(9.5), I(0.8), size=40)
for k, t in enumerate(AX):
    x = L0 + k * I(3.95)
    cd = card(s, f'!!Card{k}', x, I(2.3), I(3.7), I(3.0))
    n = text(s, i, f'!!CardN{k}', x + I(0.35), I(2.55), I(2), I(1.0), [[(f'0{k + 1}', F_TITLE, 54, BORD, {})]], row='axn')
    tt = text(s, i, f'!!CardT{k}', x + I(0.35), I(3.75), I(3.05), I(1.2), [[(t, F_SEMI, 17, TXT, {})]], row='axt')
    ag(i, k + 1, cd, n, tt)
finish(s, i, m, '01 · Axes de communication', secs=20,
       note="Trois axes guident tous les supports (p. 28) : l'authenticité des saveurs ivoiriennes, la convivialité et le partage, la simplicité accessible au quotidien.")

# ---------------------------------------------------------------- 13-18 : séquence produits, Morph 3D conservé
# 13 — Intercalaire Produits
i = 13; m = 'D'; s = new(i, m, dur=1600)
text(s, None, '!!SecNum', W / 2 - I(3), I(1.35), I(6), I(1.4), [[('02', F_TITLE, 90, OR, {})]], align='c', anchor='m')
products(s, i)
title(s, i, m, 'Présentation des produits', 0, H / 2 - I(0.45), W, I(0.9), size=44, align='c', anchor='m')
finish(s, i, m, '02 · Les produits', logo=False, secs=10, note="Passons aux produits. Les quinze références de la gamme apparaissent ensemble (p. 24).")

# 14 — Line-up
i = 14; m = 'L'; s = new(i, m, dur=2000)
rect(s, '!!Table', 0, I(5.55), W, H - I(5.55), grad=grad_xml([(0, ROSE, 0), (100, ROSE, 100)], angle=90))
products(s, i)
title(s, i, m, '15 produits, 3 familles', 0, I(3.6), W, I(0.62), size=30, align='c')
text(s, i, '!!Sub', 0, I(4.22), W, I(0.3), [[('Épices  ·  Marinades  ·  Essentiels', F_SEMI, 12, ROUGE, dict(spc=250, caps=True))]], align='c')
finish(s, i, m, '02 · La gamme', secs=20,
       note="Au moment de notre intervention, la marque disposait de quinze produits répartis en trois familles (Tableau 1, p. 24) : dix épices et condiments, trois marinades, deux essentiels de la cuisine ivoirienne.")

# 15 — Marinades
i = 15; m = 'L'; s = new(i, m, dur=1900)
rect(s, '!!Table', 0, I(6.4), W, H - I(6.4), grad=grad_xml([(0, ROSE, 0), (100, ROSE, 100)], angle=90))
products(s, i)
title(s, i, m, 'Les Marinades', L0, I(1.0), I(5.5), I(0.8), size=40)
body(s, i, m, '!!Body', ["Prêtes à l'emploi, en pot.",
                         [('Poulet', F_SEMI, None, BORD, {}), (' : braisé, rôti ou grillé', None, None, None, {})],
                         [('Poisson', F_SEMI, None, BORD, {}), (' : braisé, frit ou fumé', None, None, None, {})],
                         [('Mixte', F_SEMI, None, BORD, {}), (' : viandes et poissons', None, None, None, {})]],
     L0, I(2.25), I(5.3), I(3.2), size=18, space=10)
finish(s, i, m, '02 · Marinades · 3 produits', secs=18,
       note="Les marinades, prêtes à l'emploi, en pot : poulet (braisé, rôti ou grillé), poisson (braisé, frit ou fumé) et mixte (viandes et poissons). Source : Tableau 1, p. 24.")

# 16 — Épices
i = 16; m = 'L'; s = new(i, m, dur=1900)
products(s, i)
title(s, i, m, 'Les Épices', L0, I(1.0), I(4.2), I(0.8), size=40)
body(s, i, m, '!!Body', ["Dix produits en sachet, pour relever le goût, colorer, parfumer et lier les sauces."], L0, I(2.25), I(4.0), I(2.0), size=18)
finish(s, i, m, '02 · Épices · 10 produits', secs=18,
       note="Dix épices et condiments, majoritairement en sachet (Tableau 1) : épices spéciales poulet et poisson, paprika, piment, graines d'akpi, poudres de djoumblé, de crevettes et de poisson, purée d'ail, bouillon de poulet. Ils servent à relever le goût, colorer, parfumer et lier les sauces.")

# 17 — Essentiels
i = 17; m = 'L'; s = new(i, m, dur=1900)
rect(s, '!!Table', 0, I(6.4), W, H - I(6.4), grad=grad_xml([(0, ROSE, 0), (100, ROSE, 100)], angle=90))
products(s, i)
title(s, i, m, 'Les Essentiels', L0, I(1.0), I(5.5), I(0.8), size=40)
body(s, i, m, '!!Body', [[("Pâte d'arachide", F_SEMI, None, BORD, {}), (' : la sauce arachide', None, None, None, {})],
                         [('Huile rouge', F_SEMI, None, BORD, {}), (' : les sauces traditionnelles', None, None, None, {})]],
     L0, I(2.25), I(5.3), I(1.6), size=18, space=10)
finish(s, i, m, '02 · Essentiels · 2 produits', secs=14,
       note="Les essentiels de la cuisine ivoirienne : la pâte d'arachide pour la sauce arachide, l'huile rouge pour les sauces traditionnelles (Tableau 1).")

# 18 — Présentoir : fond totalement uni
i = 18; m = 'D'; s = new(i, m, solid=True, dur=2000)
from layout3d import TIERS
for n, (by, k, items) in enumerate(TIERS):
    xs = [it['x'] - WID[it['key']] * k / 2 for it in items] + [it['x'] + WID[it['key']] * k / 2 for it in items]
    x0, x1 = min(xs) - I(0.35), max(xs) + I(0.35)
    ty = by - I(0.05)
    rect(s, f'!!Tier{n}', x0, ty, x1 - x0, H - ty, fill='5A0E18')
    rect(s, f'!!TierEdge{n}', x0, ty, x1 - x0, I(0.025), fill=OR)
    products(s, i, only={it['key'] for it in items})
title(s, i, m, 'Le vrai goût de chez nous', 0, I(0.42), W, I(0.8), size=40, align='c')
body(s, i, m, '!!Sub', ["Marinades et épices à base d'ingrédients 100 % naturels"], 0, I(1.22), W, I(0.4), size=16, align='c', color=CREME)
finish(s, i, m, None, logo=False, secs=15,
       note="« Le vrai goût de chez nous » : des marinades et épices à base d'ingrédients 100 % naturels (accroche et sous-texte de l'affiche 1, p. 47). C'est cette promesse que l'identité visuelle doit porter.")


# ---------------------------------------------------------------- intercalaires
def divider(i, num, txt, kick, note, secs=8):
    s = new(i, 'D')
    n = text(s, None, '!!SecNum', W / 2 - I(3), I(0.95), I(6), I(1.9), [[(num, F_TITLE, 120, OR, {})]], align='c', anchor='m')
    ag(i, 0, n)
    title(s, i, 'D', txt, I(1.0), H / 2 - I(0.45), W - I(2.0), I(0.9), size=40, align='c', anchor='m')
    finish(s, i, 'D', kick, secs=secs, note=note)
    return s


divider(19, '03', 'Identité visuelle et charte graphique', '03 · Identité visuelle', "Troisième temps : l'identité visuelle.")

# 20 — Trois pistes
PIS = [('rough1.jpg', 'Piste 1 · Emblème', 'Couleurs du drapeau : Accord de Bangui'), ('rough2.jpg', 'Piste 2 · Symbole', 'Éléphant et vapeur, peu lisible'),
       ('rough3.jpg', 'Piste 3 · Retenue', 'Éléphant et feuille, la plus symbolique')]
i = 20; m = 'L'; s = new(i, m)
title(s, i, m, 'Trois pistes explorées', L0, I(1.0), I(9), I(0.8), size=40)
for k, (f, t, d) in enumerate(PIS):
    x = L0 + k * I(3.95)
    cd = card(s, f'!!Card{k}', x, I(2.2), I(3.7), I(2.75), fill=BLANC)
    set_line(cd, 'E8D3D6', 1.0)
    im = picture(s, f'!!Img{k}', A + f, x + I(0.2), I(2.4), I(3.3), I(2.35), fit='contain', descr=t)
    tt = text(s, i, f'!!CardT{k}', x, I(5.15), I(3.7), I(0.35), [[(t, F_SEMI, 15, BORD if k == 2 else TXT, {})]], row='pt')
    dd = text(s, i, f'!!CardD{k}', x, I(5.52), I(3.7), I(0.35), [[(d, F_MED, 12.5, MUTED, {})]], row='pd')
    ag(i, k + 1, cd, im, tt, dd)
finish(s, i, m, '03 · Rough', secs=22,
       note="Trois pistes (p. 34-35). La piste 1 reprenait l'ordre des couleurs du drapeau : l'Accord de Bangui interdit de reproduire les emblèmes d'un État sans autorisation. La piste 2, éléphant et lignes de vapeur, était difficile à identifier. La piste 3, éléphant et feuille, a été retenue, la plus narrative et la plus symbolique.")

# 21 — Logo retenu
i = 21; m = 'L'; s = new(i, m)
p = card(s, '!!Panel', L0, I(1.25), I(5.6), I(4.9))
lw = I(4.7); lh = lw * 669 / 1274
lg = picture(s, '!!LogoHero', A + 'logo_couleur.png', L0 + I(5.6) / 2 - lw / 2, I(1.25) + I(4.9) / 2 - lh / 2, lw, lh, fit='contain', descr="Logo Marin'Or Ivoire")
title(s, i, m, 'Le logo retenu', I(6.85), I(1.0), I(5.7), I(0.8), size=40)
b = body(s, i, m, '!!Body', [[('Éléphant', F_SEMI, None, BORD, {}), (" : la Côte d'Ivoire, la force, l'ancrage local.", None, None, None, {})],
                             [('Feuille', F_SEMI, None, BORD, {}), (' : épices, herbes, naturalité, fraîcheur.', None, None, None, {})],
                             [('Bordeaux', F_SEMI, None, BORD, {}), (' : premium, chaleureux, gourmand.', None, None, None, {})]],
         I(6.85), I(2.25), I(5.6), I(3.6), size=17, space=12)
ag(i, 1, p, lg); ag(i, 2, b)
finish(s, i, m, '03 · Logotype définitif', logo=False, secs=25,
       note="Le logo associe une tête d'éléphant de profil à une feuille qui en forme l'oreille (p. 35-37). L'éléphant renvoie au mot « Ivoire » : force, mémoire, ancrage local. La feuille évoque l'univers culinaire. Le bordeaux profond donne un caractère premium et gourmand ; le doré, noblesse et qualité.")

# 22 — Zone de protection et déclinaisons
i = 22; m = 'L'; s = new(i, m)
h = pic(s, '!!Hero', A + 'zone.jpg', L0, I(2.05), I(5.6), I(3.2), radius=0.04, descr='Zone de protection du logo')
h2 = pic(s, '!!Second', A + 'declinaisons.jpg', I(6.6), I(2.05), I(5.98), I(3.4), radius=0.04, fit='contain', descr='Déclinaisons du logo')
title(s, i, m, 'Zone de protection et déclinaisons', L0, I(1.0), I(10), I(0.8), size=36)
b = body(s, i, m, '!!Body', ['Couleur pour fonds clairs  ·  Blanc pour fonds sombres  ·  Noir pour une seule couleur'], L0, I(5.75), TW, I(0.45), size=15, align='c')
ag(i, 1, h); ag(i, 2, h2, b)
finish(s, i, m, '03 · Charte graphique', secs=15,
       note="Une zone de protection entoure le logotype. Trois versions : couleur pour les fonds clairs, blanche pour les fonds sombres et les photos, noire pour l'impression en une couleur (p. 37-38). Cette présentation applique la règle : logo couleur sur fond clair, blanc sur fond sombre.")

# 23 — Interdits
i = 23; m = 'L'; s = new(i, m)
h = pic(s, '!!Hero', A + 'interdits.jpg', I(2.6), I(1.95), I(8.1), I(3.55), radius=0.04, fit='contain', descr='Les interdits du logo')
title(s, i, m, 'Les interdits', L0, I(1.0), I(9), I(0.8), size=40)
b = body(s, i, m, '!!Body', ["Les manipulations à éviter pour préserver l'intégrité du logotype."], L0, I(5.85), TW, I(0.45), size=15, align='c')
ag(i, 1, h, b)
finish(s, i, m, '03 · Charte graphique', secs=10, note="Certaines manipulations sont proscrites pour préserver l'intégrité et la reconnaissance du logotype (p. 38).")

# 24 — Typographie
TYP = [(F_SEMI, 'Montserrat', 'Titres, sous-titres et textes'), (F_TITLE, 'Cream Heavy', 'Accroches et titres forts'), (None, 'RadioTime Tooled', 'Réservée au logotype')]
i = 24; m = 'L'; s = new(i, m)
title(s, i, m, 'Typographie', L0, I(1.0), I(9), I(0.8), size=40)
for k, (f, n, d) in enumerate(TYP):
    x = L0 + k * I(3.95)
    cd = card(s, f'!!Card{k}', x, I(2.2), I(3.7), I(3.4))
    if f:
        sp = text(s, i, f'!!Aa{k}', x + I(0.35), I(2.4), I(3.0), I(1.6), [[('Aa', f, 72, BORD, {})]], row='aa')
    else:
        lw = I(2.6); lh = lw * 669 / 1274
        sp = picture(s, '!!Spec', A + 'logo_noir.png', x + I(0.55), I(2.55), lw, lh, fit='contain', descr='Wordmark en RadioTime Tooled')
    tt = text(s, i, f'!!CardT{k}', x + I(0.35), I(4.15), I(3.1), I(0.4), [[(n, F_SEMI, 16, TXT, {})]], row='tn')
    dd = text(s, i, f'!!CardD{k}', x + I(0.35), I(4.6), I(3.1), I(0.4), [[(d, F_MED, 13, MUTED, {})]], row='td')
    ag(i, k + 1, cd, sp, tt, dd)
finish(s, i, m, '03 · Charte graphique', secs=18,
       note="Trois familles (p. 39) : Montserrat pour la lisibilité sur tous les supports ; Cream Heavy, généreuse et gourmande, pour les accroches ; RadioTime Tooled, calligraphique, réservée au logotype. Cette présentation utilise elle-même Cream Heavy et Montserrat.")

# 25 — Colorimétrie
COL = [('Bordeaux Profond', BORD), ('Rouge Passion', ROUGE), ('Or Épicé', OR), ('Jaune Soleil', JAUNE), ('Vert Nature', VERT),
       ('Vert Vif', VERTV), ('Crème Douce', CREME), ('Rose Poudre', ROSE), ('Blanc Pureté', BLANC), ('Noir Intense', NOIR)]
i = 25; m = 'L'; s = new(i, m)
title(s, i, m, 'Colorimétrie', L0, I(1.0), I(9), I(0.8), size=40)
for k, (n, c) in enumerate(COL):
    x = L0 + (k % 5) * I(2.4); y = I(2.15) + (k // 5) * I(2.15)
    sw = rect(s, f'!!Sw{k}', x, y, I(2.15), I(1.25), fill=c, geom=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.1,
              line='E3CED1' if c in (BLANC, CREME, ROSE) else None, line_w=0.75)
    a = text(s, i, f'!!SwN{k}', x, y + I(1.35), I(2.2), I(0.3), [[(n, F_SEMI, 12.5, TXT, {})]], row=f'swn{k // 5}')
    b = text(s, i, f'!!SwH{k}', x, y + I(1.64), I(2.2), I(0.28), [[('#' + c, F_MED, 11, MUTED, dict(spc=80))]], row=f'swh{k // 5}')
    ag(i, 1 + k // 5, sw, a, b)
finish(s, i, m, '03 · Charte graphique', secs=18,
       note="Dix couleurs (p. 40). Le bordeaux profond, couleur du logotype ; le rouge passion ; l'or épicé et le jaune soleil, gourmands ; les deux verts pour la naturalité ; la crème douce et le rose poudre pour les respirations ; le noir et le blanc pour la lisibilité.")

# 26 — Mise en application
MOCK = [('carte.jpg', 'Carte de visite'), ('carnet.jpg', 'Carnet'), ('polos.jpg', 'Polos'), ('tabliers.jpg', 'Tabliers'),
        ('camion.jpg', 'Camion de livraison'), ('sac.jpg', 'Sac kraft')]
i = 26; m = 'L'; s = new(i, m)
title(s, i, m, "Mise en application", L0, I(1.0), I(9), I(0.8), size=40)
gw, gh, gx, gy = I(3.78), I(2.1), I(0.17), I(0.2)
for k, (f, d) in enumerate(MOCK):
    x = L0 + (k % 3) * (gw + gx); y = I(2.0) + (k // 3) * (gh + gy)
    ag(i, 1 + k // 3, pic(s, f'!!M{k}', A + f, x, y, gw, gh, radius=0.05, descr=d))
finish(s, i, m, '03 · Identité appliquée', secs=15,
       note="Pour vérifier que l'identité fonctionne en situation réelle (p. 41-44), elle a été appliquée à une carte de visite, un carnet, des polos, des tabliers de cuisine, un camion de livraison et un sac kraft.")

# 27 — Animation du logo
i = 27; m = 'L'; s = new(i, m)
vx, vy, vw = I(5.6), I(1.65), I(7.0); vh = vw * 9 / 16
mv = s.shapes.add_movie(VID + 'logo_anim.mp4', Emu(vx), Emu(vy), Emu(vw), Emu(int(vh)), poster_frame_image=A + 'logo_end.jpg', mime_type='video/mp4')
mv.name = '!!Video'
title(s, i, m, 'Animation du logotype', L0, I(1.0), I(4.6), I(1.6), size=36)
b = body(s, i, m, '!!Body', ['La signature audiovisuelle qui ouvre et clôt les spots.', 'Illustrator, puis After Effects · Full HD'], L0, I(2.75), I(4.5), I(2.6), size=17, space=12)
ag(i, 1, b)
finish(s, i, m, '03 · Motion design', secs=12,
       note="Cliquer sur la vidéo pour la lancer. L'animation du logotype est la signature audiovisuelle de la marque : elle ouvre ou clôt les spots et vidéos (p. 44). Préparation sous Illustrator, animation sous After Effects.")

divider(28, '04', 'Supports print et audiovisuels', '04 · Réalisations', "Quatrième temps : les supports réalisés.")

# 29-31 — Affiches
AFF = [('affiche1', 'Le vrai goût de chez nous', "Axe 1 · L'authenticité des saveurs ivoiriennes", 18,
        "Affiche 1 (p. 47, 86) : accroche « Le vrai goût de chez nous », sous-texte « Marinades et épices à base d'ingrédients 100 % naturels », un jeune homme goûtant une sauce, trois produits mis en avant, un QR code et la barre de contacts. Axe : authenticité."),
       ('affiche2', 'Le goût qui rassemble', 'Axe 2 · La convivialité et le partage', 15,
        "Affiche 2 (p. 47, 87) : « Le goût qui rassemble », « Invitez Marin'Or Ivoire à votre table ! », des mains se servant dans un grand plat traditionnel, les trois marinades. Axe : convivialité et partage."),
       ('affiche3', 'Marinez, Cuisez, Régalez-vous', 'Axe 3 · La simplicité accessible au quotidien', 15,
        "Affiche 3 (p. 47, 89) : « Marinez, Cuisez, Régalez-vous », « Aussi facile que de former une boule d'attiéké ! », deux produits. Axe : simplicité au quotidien.")]
for n, (f, t, ax, secs, nt) in enumerate(AFF):
    i = 29 + n; m = 'L'; s = new(i, m)
    h = pic(s, '!!Hero', A + f + '.jpg', I(4.85), I(1.45), I(7.75), I(7.75) * 6087 / 10304, radius=0.025, descr=f'Affiche {n + 1}')
    h2 = pic(s, '!!Second', A + f + '_situ.jpg', L0, I(3.85), I(3.75), I(2.75), radius=0.04, descr=f'Affiche {n + 1} en situation')
    title(s, i, m, t, L0, I(1.0), I(3.85), I(1.7), size=30)
    a = text(s, i, '!!Axe', L0, I(2.85), I(3.85), I(0.7), [[(ax, F_SEMI, 14, ROUGE, {})]])
    ag(i, 1, a, h); ag(i, 2, h2)
    finish(s, i, m, f'04 · Affiche {n + 1}', secs=secs, note=nt)

# 32 — Dépliant
i = 32; m = 'L'; s = new(i, m)
h = pic(s, '!!Hero', A + 'depliant_recto.jpg', I(4.85), I(1.15), I(5.0), I(5.0) * 623 / 871, radius=0.02, descr='Dépliant recto (extérieur)')
h3 = pic(s, '!!Third', A + 'depliant_verso.jpg', I(7.6), I(3.45), I(5.0), I(5.0) * 625 / 872, radius=0.02, shadow=(18, 4, 18), descr='Dépliant verso (intérieur)')
h2 = pic(s, '!!Second', A + 'depliant_situ.jpg', L0, I(3.85), I(3.75), I(2.75), radius=0.04, descr='Dépliant en situation')
title(s, i, m, 'Dépliant trois volets', L0, I(1.0), I(3.85), I(1.7), size=30)
a = text(s, i, '!!Axe', L0, I(2.85), I(3.85), I(0.7), [[('Les trois axes · extérieur : la marque, intérieur : la gamme', F_SEMI, 14, ROUGE, {})]])
ag(i, 1, a, h); ag(i, 2, h3); ag(i, 3, h2)
finish(s, i, m, '04 · Dépliant', secs=20,
       note="Dépliant A4 trois volets (p. 47-48) : à l'extérieur, la marque, « Qui sommes-nous ? », « Pourquoi nous choisir ? » et les contacts ; à l'intérieur, la gamme en deux colonnes, Marinades et Épices & Essentiels. Il porte les trois axes.")

# 33 — Flyer
i = 33; m = 'L'; s = new(i, m)
fh = I(5.45); fw = fh * 5000 / 7114
h = pic(s, '!!Hero', A + 'flyer.jpg', I(8.45), I(1.15), fw, fh, radius=0.02, descr='Flyer')
h2 = pic(s, '!!Second', A + 'flyer_situ.jpg', I(4.85), I(1.6), I(3.3), I(3.3) * 1445 / 1088, radius=0.04, descr='Flyer en situation')
title(s, i, m, 'Marinez, Cuisez, Régalez-vous !', L0, I(1.0), I(3.85), I(1.7), size=30)
a = text(s, i, '!!Axe', L0, I(2.85), I(3.85), I(0.7), [[("Flyer · Axe 3 · « Le goût du chef, sans l'effort »", F_SEMI, 14, ROUGE, {})]])
ag(i, 1, a, h); ag(i, 2, h2)
finish(s, i, m, '04 · Flyer', secs=15,
       note="Le flyer (p. 48, 93) reprend « Marinez, Cuisez, Régalez-vous ! » et « Le goût du chef, sans l'effort », avec une sélection de produits alignés en bas. Axe : simplicité au quotidien.")

# 34 — Kakémono
i = 34; m = 'L'; s = new(i, m)
kh = I(5.6); kw = kh * 5000 / 12500
h = pic(s, '!!Hero', A + 'kakemono.jpg', I(10.3), I(1.1), kw, kh, radius=0.03, descr='Kakémono')
h2 = pic(s, '!!Second', A + 'kakemono_situ.jpg', I(4.85), I(1.6), I(5.1), I(5.1) * 1404 / 1532, radius=0.04, descr='Kakémono en situation')
title(s, i, m, 'Le goût qui rassemble !', L0, I(1.0), I(3.85), I(1.7), size=30)
a = text(s, i, '!!Axe', L0, I(2.85), I(3.85), I(0.7), [[('Kakémono · Axe 2 · point de vente et événements', F_SEMI, 14, ROUGE, {})]])
ag(i, 1, a, h); ag(i, 2, h2)
finish(s, i, m, '04 · Kakémono', secs=15,
       note="Le kakémono (p. 48-49, 96) sert en point de vente et lors d'événements : « Le goût qui rassemble ! », « Profitez-en ! », le plat partagé et les trois gammes. Axe : convivialité et partage.")

# 35-36 — Spots
SP = [('spot2d_clip.mp4', 'spot2d_poster.jpg', 'Spot 2D', ['Une question, des solutions naturelles, la gamme complète, puis le repas en famille.', 'After Effects · 1920 × 1080'], 25,
       "Cliquer sur la vidéo pour la lancer (extrait de 14 s). Le spot 2D (p. 49-50, 98) part d'une question au spectateur, présente les solutions naturelles Marin'Or, anime toute la gamme, illustre la simplicité d'usage et se clôt sur un repas en famille et le slogan « Marin'Or Ivoire, le goût en or »."),
      ('film3d_clip.mp4', 'film3d_poster.jpg', 'Spot 3D', ["Une visite immersive de la boutique virtuelle Marin'Or Ivoire, modélisée sous Blender."], 25,
       "Cliquer sur la vidéo pour la lancer (extrait de 15 s). Le spot 3D (p. 50-51, 99) propose une visite de la boutique virtuelle : la caméra entre, découvre les rayonnages, s'attarde sur les marinades, les épices et les essentiels, puis termine sur le logo et l'appel à l'action.")]
for n, (v, pst, t, bd, secs, nt) in enumerate(SP):
    i = 35 + n; m = 'L'; s = new(i, m)
    vx, vy, vw = I(5.1), I(1.55), I(7.5); vh = vw * 9 / 16
    mv = s.shapes.add_movie(VID + v, Emu(vx), Emu(vy), Emu(vw), Emu(int(vh)), poster_frame_image=A + pst, mime_type='video/mp4')
    mv.name = '!!Video'
    title(s, i, m, t, L0, I(1.0), I(4.0), I(0.8), size=40)
    b = body(s, i, m, '!!Body', bd, L0, I(2.25), I(4.0), I(3.5), size=17, space=12)
    ag(i, 1, b)
    finish(s, i, m, '04 · Audiovisuel', secs=secs, note=nt)

divider(37, '05', 'Le site web : une boutique en ligne', '05 · Site web', "Cinquième temps : le site web.")

# 38 — Structure de navigation
i = 38; m = 'L'; s = new(i, m)
title(s, i, m, 'Structure de navigation', L0, I(1.0), I(9), I(0.8), size=40)
for k, (f, lab) in enumerate([('arbo_client.jpg', 'Partie clientèle'), ('arbo_admin.jpg', 'Back-office')]):
    x = L0 + k * I(6.05)
    cd = card(s, f'!!Card{k}', x, I(2.1), I(5.75), I(3.65))
    im = picture(s, f'!!Img{k}', A + f, x + I(0.2), I(2.3), I(5.35), I(3.25), fit='contain', descr=lab)
    tt = text(s, i, f'!!CardT{k}', x, I(5.95), I(5.75), I(0.4), [[(lab, F_SEMI, 15, BORD, {})]], align='c', row='arb')
    ag(i, k + 1, cd, im, tt)
finish(s, i, m, '05 · Conception', secs=20,
       note="Deux espaces (p. 52-53) : la partie clientèle, ouverte à tout visiteur (boutique, gammes, fiches, panier, tunnel de commande, compte), et le back-office réservé à l'administrateur, avec un tableau de bord et six rubriques de gestion.")

# 39 — Accueil et gammes
i = 39; m = 'L'; s = new(i, m)
h = pic(s, '!!Hero', A + 'site_00.jpg', I(5.6), I(1.95), I(7.0), I(7.0) * 0.5, radius=0.02, descr="Page d'accueil")
ths = [pic(s, f'!!Th{k}', A + f, I(5.6) + k * I(2.4), I(5.58), I(2.2), I(1.05), radius=0.04, descr=['Gamme Marinades', 'Gamme Épices', 'Gamme Essentiels'][k])
       for k, f in enumerate(['site_02.jpg', 'site_03.jpg', 'site_04.jpg'])]
title(s, i, m, 'Accueil et gammes', L0, I(1.0), I(9), I(0.8), size=40)
b = body(s, i, m, '!!Body', ['Une couleur par gamme : bordeaux, or, vert.'], L0, I(2.25), I(4.3), I(3.0), size=16.5, space=12)
ag(i, 1, h); ag(i, 2, b, *ths)
finish(s, i, m, '05 · Interfaces', secs=20,
       note="L'accueil s'ouvre sur une bannière bordeaux avec l'accroche « Le vrai goût de chez nous » (p. 101). Chaque gamme a sa couleur de charte : bordeaux pour les Marinades, or pour les Épices, vert pour les Essentiels. Filtres, recherche et tri sont disponibles.")

# 40 — Commander et payer
i = 40; m = 'L'; s = new(i, m)
h = pic(s, '!!Hero', A + 'site_10.jpg', I(5.6), I(1.95), I(7.0), I(7.0) * 0.5, radius=0.02, descr='Livraison et paiement')
t0 = pic(s, '!!Th0', A + 'site_07.jpg', I(5.6), I(5.58), I(3.4), I(1.05), radius=0.04, descr='Panier')
t1 = pic(s, '!!Th1', A + 'site_11.jpg', I(9.2), I(5.58), I(3.4), I(1.05), radius=0.04, descr='Paiement et confirmation')
title(s, i, m, 'Commander et payer', L0, I(1.0), I(9), I(0.8), size=40)
b = body(s, i, m, '!!Body', ["Paiement mobile money, vérifié par l'administrateur."], L0, I(2.25), I(4.3), I(1.5), size=16.5)
ag(i, 1, h, b); ag(i, 3, t0, t1)
for k, t in enumerate(['Wave', 'Orange Money', 'Moov Money', 'Djamo']):
    x = L0 + (k % 2) * I(1.95); y = I(4.35) + (k // 2) * I(0.62)
    pl = rect(s, f'!!Pill{k}', x, y, I(1.82), I(0.48), fill=ROSE, geom=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
    pt = text(s, i, f'!!PillT{k}', x, y, I(1.82), I(0.48), [[(t, F_SEMI, 12, BORD, {})]], align='c', anchor='m', row=f'mm{k // 2}')
    ag(i, 2, pl, pt)
finish(s, i, m, '05 · Fonctionnalités', secs=25,
       note="Le tunnel de commande compte trois étapes (p. 65-66). Le client choisit Wave, Orange Money, Moov Money ou Djamo, effectue le transfert, saisit le numéro payeur et la référence. Le paiement passe « à vérifier » ; l'administrateur le valide et la commande devient « payée ». Il n'y a pas de paiement automatique : il faudrait un compte marchand, mais la base est prête pour cette évolution.")

# 41 — Back-office
i = 41; m = 'L'; s = new(i, m)
h = pic(s, '!!Hero', A + 'site_16.jpg', I(5.6), I(1.95), I(7.0), I(7.0) * 0.5, radius=0.02, descr="Tableau de bord de l'administrateur")
t0 = pic(s, '!!Th0', A + 'site_17.jpg', I(5.6), I(5.58), I(3.4), I(1.05), radius=0.04, descr='Liste des commandes')
t1 = pic(s, '!!Th1', A + 'site_18.jpg', I(9.2), I(5.58), I(3.4), I(1.05), radius=0.04, descr='Liste des produits')
title(s, i, m, 'Le back-office', L0, I(1.0), I(9), I(0.8), size=40)
b = body(s, i, m, '!!Body', ['Paiements à vérifier, commandes à préparer, ventes encaissées, clients inscrits.'], L0, I(2.25), I(4.3), I(2.5), size=16.5)
ag(i, 1, h, b); ag(i, 2, t0, t1)
finish(s, i, m, '05 · Administration', secs=18,
       note="Le back-office (p. 106-110) : un tableau de bord à quatre indicateurs, la gestion des commandes, des produits, des gammes, des clients, des messages et des paramètres. Le logo envoyé depuis les paramètres s'affiche sur tout le site.")

# 42 — Base de données
i = 42; m = 'L'; s = new(i, m)
cd = card(s, '!!Card0', I(5.3), I(1.95), I(7.3), I(4.65), fill=BLANC)
set_line(cd, 'E8D3D6', 1.0)
h = picture(s, '!!Hero', A + 'mcd.jpg', I(5.45), I(2.05), I(7.0), I(4.45), fit='contain', descr='Modèle conceptuel de données')
title(s, i, m, 'Base de données', L0, I(1.0), I(9), I(0.8), size=40)
st = [text(s, i, f'!!Stat{k}', L0, I(2.25) + k * I(1.15), I(4.3), I(1.0), [[(n + ' ', F_TITLE, 48, BORD, {}), (t, F_SEMI, 17, TXT, {})]], anchor='m')
      for k, (n, t) in enumerate([('10', 'tables'), ('11', 'règles de gestion')])]
b = body(s, i, m, '!!Body', ['Méthode MERISE · MySQL'], L0, I(4.65), I(4.0), I(0.5), size=16.5, color=MUTED)
ag(i, 1, cd, h); ag(i, 2, st[0]); ag(i, 3, st[1], b)
finish(s, i, m, '05 · Base de données', secs=20,
       note="La base marinor_db a été conçue avec MERISE (p. 68-83) : MCD, MLD puis MPD. Onze règles de gestion, un dictionnaire de 74 données, dix entités devenues dix tables MySQL (moteur InnoDB, encodage utf8mb4). Sécurité : requêtes préparées, mots de passe hachés, jetons CSRF.")

# 43 — Tests
i = 43; m = 'L'; s = new(i, m)
TST = [('14', 'conformes', VERT), ('2', 'à vérifier', BORD), ('1', 'non conforme', ROUGE)]
title(s, i, m, 'Tests fonctionnels', L0, I(1.0), I(9), I(0.8), size=40)
for k, (n, t, c) in enumerate(TST):
    x = L0 + k * I(3.95)
    cd = card(s, f'!!Card{k}', x, I(2.15), I(3.7), I(2.55))
    sn = text(s, i, f'!!StatN{k}', x, I(2.3), I(3.7), I(1.3), [[(n, F_TITLE, 72, c, {})]], align='c', anchor='m', row='tn')
    stt = text(s, i, f'!!StatT{k}', x, I(3.75), I(3.7), I(0.45), [[(t, F_SEMI, 17, TXT, {})]], align='c', row='tt')
    ag(i, k + 1, cd, sn, stt)
b = body(s, i, m, '!!Body', ['17 tests en local, parcours client puis administrateur. Non conforme : affichage sur smartphone.'], L0, I(5.1), TW, I(0.9), size=15.5, align='c', color=MUTED)
ag(i, 4, b)
finish(s, i, m, '05 · Tests', secs=25,
       note="Dix-sept tests en local (Tableau 16, p. 111) : 14 conformes. Deux restent « à vérifier » : la saisie et la validation du paiement. Un est non conforme : l'affichage sur smartphone, à corriger avant la mise en ligne. ATTENTION : vérifier ces deux résultats avant la soutenance.")

# ---------------------------------------------------------------- 44 — Calendrier : barres segmentées (inspiration graph2)
CAL = [('Identité visuelle', [('Logotype', 2), ('Charte graphique', 7)]),
       ('Supports imprimés', [('Affiches', 3), ('Flyer', 1), ('Dépliant', 1), ('Kakémono', 1)]),
       ('Spots animatiques', [('Scénario et synopsis', 3), ('Storyboard', 2), ('Production', 10), ('Post-production', 5)]),
       ('Site web', [('Recherche de données', 10), ('Maquette', 4), ('Interfaces', 4), ('Base de données', 3), ('Tests', 3), ('Finalisation', 4)])]
i = 44; m = 'L'; s = new(i, m)
t = text(s, i, '!!Title', L0, I(0.95), I(10), I(0.95), [[('63 jours', F_TITLE, 46, BORD, {}), ('   soit 9 semaines de réalisation', F_MED, 16, MUTED, {})]], anchor='b')
ag(i, 0, t)
X0, BW, RH, Y0 = I(3.55), I(8.1), I(0.98), I(2.45)
ppd = BW / 28                                     # EMU per day (longest bar = 28 days)
for wk in range(0, 5):                            # weekly grid: 0, 7, 14, 21, 28 days
    gx_ = X0 + wk * 7 * ppd
    rect(s, f'!!Grid{wk}', gx_, Y0 - I(0.25), I(0.01), 4 * RH + I(0.1), fill=BORD, alpha=12)
chart_steps = []
for r, (lab, segs) in enumerate(CAL):
    y = Y0 + r * RH
    tl = text(s, i, f'!!BarL{r}', L0, y, I(2.7), I(0.5), [[(lab, F_SEMI, 15, TXT, {})]], anchor='m', row=f'cal{r}')
    x = X0
    row_steps = []
    for k, (sn_, d) in enumerate(segs):
        w = d * ppd - I(0.035)
        sg = rect(s, f'!!Seg{r}_{k}', x, y + I(0.06), w, I(0.38), fill=BORD if k % 2 == 0 else TINT,
                  geom=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.35)
        sg.name = f'!!Seg{r}_{k}'
        row_steps.append([('wipe', sg.shape_id, max(160, 34 * d))])
        x += d * ppd
    tot = sum(d for _, d in segs)
    tv = text(s, i, f'!!BarV{r}', x + I(0.12), y, I(1.0), I(0.5), [[(f'{tot} j', F_TITLE, 20, BORD, {})]], anchor='m', row=f'cal{r}')
    chart_steps += [[('float', tl.shape_id, 400)]] + row_steps + [[('fade', tv.shape_id, 300)]]
SL[i]['chart'] = [chart_steps]
finish(s, i, m, '06 · Calendrier d’exécution', secs=22,
       note="Le projet a demandé 63 jours, soit 9 semaines (Tableau 2, p. 31). Chaque barre se construit tâche par tâche : identité visuelle 9 jours (logotype 2, charte 7) ; supports imprimés 6 jours ; spots 20 jours (scénario 3, storyboard 2, production 10, post-production 5) ; site web 28 jours (recherche 10, maquette 4, interfaces 4, base de données 3, tests 3, finalisation 4). Les traits verticaux marquent les semaines.")

# ---------------------------------------------------------------- 45 — Coût : camembert « verre dépoli » sur halo (inspiration graph1)
CO = [('Logo et identité visuelle', 500000), ('Supports imprimés', 600000), ('Spots 2D et 3D', 2400000), ('Site web', 3500000)]
TOT = sum(v for _, v in CO)
i = 45; m = 'L'; s = new(i, m)
title(s, i, m, 'Coût estimatif : 7 000 000 FCFA', L0, I(1.0), I(11), I(0.8), size=40)
cx, cy, R = I(4.0), I(4.45), I(1.95)
panel = rect(s, '!!ChartPanel', L0, I(2.0), I(6.55), I(4.9), fill=ROSE, geom=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
orb_d = I(3.5)
orb = rect(s, '!!Orb', cx - orb_d / 2, cy - orb_d / 2, orb_d, orb_d, geom=MSO_SHAPE.OVAL,
           grad=grad_xml([(0, ROUGE, None), (38, ROUGE, 15), (68, BORD, 55), (100, BORD, 100)], path=(50000, 50000, 50000, 50000)))
fmt = lambda v: f'{v:,}'.replace(',', ' ')
ang, gap = -90.0, 2.4
steps_w, steps_l = [], []
lab_x = I(7.35)
for k, (lab, v) in enumerate(CO):
    sweep = 360 * v / TOT
    a0, a1 = ang + gap / 2, ang + sweep - gap / 2
    wd = rect(s, f'!!Wedge{k}', cx - R, cy - R, 2 * R, 2 * R, geom=MSO_SHAPE.PIE, fill=BLANC, alpha=42,
              line=BLANC, line_w=1.5, shadow=(26, 5, 20))
    av = wd._element.spPr.find('a:prstGeom/a:avLst', NS)
    for e in list(av):
        av.remove(e)
    av.append(X(f'<a:gd name="adj1" fmla="val {int((a0 % 360) * 60000)}"/>'))
    av.append(X(f'<a:gd name="adj2" fmla="val {int((a1 % 360) * 60000)}"/>'))
    steps_w.append([('wheel', wd.shape_id, 520)])
    # leader line: from the wedge edge (mid-angle; bottom edge for the large left slice) out to the label column
    mid = math.radians((a0 + a1) / 2 if k < 3 else 93)
    px_, py_ = cx + R * math.cos(mid), cy + R * math.sin(mid)
    ex, ey = cx + (R + I(0.2)) * math.cos(mid), cy + (R + I(0.2)) * math.sin(mid)
    fb = s.shapes.build_freeform(int(px_), int(py_), scale=1.0)
    fb.add_line_segments([(int(ex), int(ey)), (int(lab_x - I(0.12)), int(ey))], close=False)
    ln = fb.convert_to_shape(); ln.name = f'!!Lead{k}'
    set_fill(ln, '<a:noFill/>'); set_line(ln, MUTED, 0.9, 60)
    st_ = ln._element.find('p:style', NS)
    if st_ is not None:
        ln._element.remove(st_)
    pct = round(100 * v / TOT)
    lt = text(s, i, f'!!Lab{k}', lab_x, ey - I(0.24), I(5.3), I(0.48),
              [[(f'{pct} %  ', F_TITLE, 21, BORD, {}), (lab, F_SEMI, 14, TXT, {}), (f'  ·  {fmt(v)}', F_MED, 14, MUTED, {})]], anchor='m')
    steps_l.append([('wipe', ln.shape_id, 300), ('float', lt.shape_id, 400)])
    ang += sweep
ag(i, 0, panel)
SL[i]['chart'] = [[[('fade', orb.shape_id, 450)]] + steps_w + steps_l]
finish(s, i, m, '06 · Coût du projet', secs=22,
       note="Le coût est estimé à 7 000 000 FCFA (Tableau 3, p. 32) : site web 3 500 000 ; spot 3D 1 500 000 et spot 2D 900 000 ; logo et identité 500 000 ; affiches 300 000, flyer, dépliant et kakémono 100 000 chacun. Les regroupements et les pourcentages (50, 34, 9 et 7 %) sont calculés à partir de ces lignes.")

# 46 — Conclusion
i = 46; m = 'L'; s = new(i, m)
h = pic(s, '!!Hero', A + 'spot2d_still.jpg', I(7.6), I(1.0), I(5.0), I(5.55), radius=0.05, descr='La gamme complète, spot 2D')
title(s, i, m, 'Conclusion', L0, I(1.0), I(6.5), I(0.8), size=40)
b1 = body(s, i, m, '!!Body', ["Une identité cohérente, des supports print et audiovisuels adaptés, un site web fonctionnel."], L0, I(2.25), I(6.5), I(1.6), size=17.5)
b2 = body(s, i, m, '!!Body2', [[('Limites', F_SEMI, None, BORD, {}), (" : pas d'étude de marché à grande échelle ; impact réel à mesurer.", None, None, None, {})]],
          L0, I(3.7), I(6.5), I(1.6), size=17.5)
ag(i, 1, b1, h); ag(i, 2, b2)
finish(s, i, m, '06 · Bilan', secs=35,
       note="Ce travail aboutit à une identité visuelle cohérente, à des supports imprimés et audiovisuels adaptés au public cible et à un site web fonctionnel, appuyé sur un modèle de données MERISE (p. 112). Limites : pas d'étude de marché approfondie à grande échelle, des choix graphiques à confronter à un panel plus large, un impact réel à mesurer une fois les supports diffusés.")

# 47 — Perspectives et merci
i = 47; m = 'D'; s = new(i, m)
lw = I(3.6); lh = lw * 669 / 1274
lg = picture(s, '!!LogoHero', A + 'logo_blanc.png', W / 2 - lw / 2, I(0.85), lw, lh, fit='contain', descr="Logo Marin'Or Ivoire")
title(s, i, m, 'Merci de votre attention', 0, I(3.05), W, I(0.9), size=48, align='c')
ag(i, 0, lg)
b = body(s, i, m, '!!Sub', ["Perspectives : une stratégie digitale (réseaux sociaux, campagnes ciblées) et l'export vers la sous-région ouest-africaine."],
         I(2.2), I(4.25), W - I(4.4), I(1.0), size=16, align='c', color=CREME)
ag(i, 1, b)
finish(s, i, m, None, logo=False, secs=15,
       note="En ouverture (p. 113) : une véritable stratégie de communication digitale, des campagnes ciblées et l'export vers la sous-région ouest-africaine. Je vous remercie de votre attention et me tiens à votre disposition pour vos questions.")

# ---------------------------------------------------------------- transitions + animations
for n, d in SL.items():
    s = d['slide']
    if n in PRODUCT_SLIDES:
        morph(s, dur=d['morph'])            # Morph only on the 3D product sequence
        continue
    g = {}
    for spid, (grp, sh) in d['groups'].items():
        g.setdefault(grp, []).append((sh.shape_id, sh))
    auto = [[('float', sid, 500) for sid, _ in g.get(0, [])]] if g.get(0) else []
    clicks = [[[('float', sid, 500) for sid, _ in g[k]]] for k in sorted(k for k in g if k > 0)]
    if d['chart']:
        clicks += d['chart']
    anim_ids = {e[1] for grp in ([auto] if auto else []) + clicks for st in grp for e in st}
    bld = []
    for sp in s.shapes._spTree.iterchildren():
        tag = etree.QName(sp).localname
        if tag == 'sp':
            sid = int(sp.find('p:nvSpPr/p:cNvPr', NS).get('id'))
            if sid in anim_ids:
                is_tb = sp.find('p:nvSpPr/p:cNvSpPr', NS).get('txBox') == '1'
                bld.append((sid, not is_tb))
    build_timing(s, auto, clicks, bld)
    SL[n]['clicks'] = len(clicks)

from pptx.opc.constants import RELATIONSHIP_TYPE as RT
th = prs.slide_master.part.part_related_by(RT.THEME)
x = th.blob.decode('utf8')
for tag, val in [('dk1', NOIR), ('lt1', BLANC), ('dk2', BORD), ('lt2', ROSE), ('accent1', BORD), ('accent2', OR), ('accent3', VERT),
                 ('accent4', ROUGE), ('accent5', JAUNE), ('accent6', VERTV)]:
    x = re.sub(rf'<a:{tag}>.*?</a:{tag}>', f'<a:{tag}><a:srgbClr val="{val}"/></a:{tag}>', x, flags=re.S)
x = re.sub(r'(<a:majorFont>\s*<a:latin typeface=")[^"]*"', r'\1Cream Heavy"', x)
x = re.sub(r'(<a:minorFont>\s*<a:latin typeface=")[^"]*"', r'\1Montserrat Medium"', x)
x = re.sub(r'<a:clrScheme name="[^"]*"', '<a:clrScheme name="Marin\'Or Ivoire"', x)
th._blob = x.encode('utf8')
cp = prs.core_properties
cp.title = "Conception et réalisation de supports de communication pour la promotion de Marin'Or Ivoire"
cp.author = 'Ekien Kouadio Constant Ephrem'; cp.subject = 'Soutenance de mémoire — Licence professionnelle Arts et Images Numériques'
os.makedirs(os.path.dirname(OUT), exist_ok=True)
prs.save(OUT)
json.dump({n - 1: dict(mode=d['mode'], secs=d.get('secs'), clicks=d.get('clicks', 0), morph=bool(d['morph'])) for n, d in SL.items()},
          open(B + '/out/slides.json', 'w'))
json.dump(lib.TEXT_LOG, open(B + '/out/textlog.json', 'w'))
print('saved', OUT, len(prs.slides), 'slides', sum(d.get('secs', 0) for d in SL.values()), 's',
      sum(d.get('clicks', 0) for d in SL.values()), 'clicks')
